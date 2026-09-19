import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

NAVY = RGBColor(31, 73, 125)
SLATE = RGBColor(54, 96, 146)
CHARCOAL = RGBColor(51, 51, 51)
LIGHT_GRAY = "F2F4F7"
BORDER_GRAY = "CCCCCC"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def apply_table_styles(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{BORDER_GRAY}"/>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="1F497D"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    # Format header row
    hdr_cells = table.rows[0].cells
    for i, cell in enumerate(hdr_cells):
        set_cell_background(cell, "1F497D")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    # Format data rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = LIGHT_GRAY if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = CHARCOAL

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                if i < len(row.cells):
                    row.cells[i].width = Inches(w)

def create_document():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / footer setup
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Pharmaceutical Research Assistant Internship | Research Portfolio")
        hrun.font.name = 'Times New Roman'
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = RGBColor(128, 128, 128)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Confidential Academic Research Portfolio | AntiGravity Enterprise Framework")
        frun.font.name = 'Times New Roman'
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(128, 128, 128)

    # Base style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = CHARCOAL
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    return doc

def add_title_page(doc, title, subtitle, author, affiliation, date_str, week_tag):
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(36)
    
    p_week = doc.add_paragraph()
    r_week = p_week.add_run(f"DELIVERABLE: {week_tag.upper()}")
    r_week.font.name = 'Times New Roman'
    r_week.font.size = Pt(12)
    r_week.font.bold = True
    r_week.font.color.rgb = SLATE
    p_week.paragraph_format.space_after = Pt(12)

    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.line_spacing = 1.15

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(14)
    r_sub.font.italic = True
    r_sub.font.color.rgb = CHARCOAL
    p_sub.paragraph_format.space_after = Pt(36)

    # Decorative divider
    p_div = doc.add_paragraph()
    r_div = p_div.add_run("―" * 45)
    r_div.font.color.rgb = SLATE
    p_div.paragraph_format.space_after = Pt(36)

    p_meta = doc.add_paragraph()
    meta_runs = [
        ("Researcher: ", True), (f"{author}\n", False),
        ("Role: ", True), ("Pharmaceutical Research Assistant\n", False),
        ("Institution / Internship: ", True), (f"{affiliation}\n", False),
        ("Date: ", True), (f"{date_str}\n", False),
        ("Framework: ", True), ("AntiGravity Enterprise Framework v2.0\n", False),
        ("Citation Style: ", True), ("American Psychological Association (APA) 7th Edition\n", False)
    ]
    for text, bold in meta_runs:
        r = p_meta.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        r.font.bold = bold
        r.font.color.rgb = CHARCOAL
    
    p_meta.paragraph_format.space_after = Pt(48)
    doc.add_page_break()

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = NAVY
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = SLATE
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = CHARCOAL
    return p

def add_callout(doc, text, alert_type="NOTE"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    
    tag = p.add_run(f"[{alert_type.upper()}] ")
    tag.font.bold = True
    tag.font.color.rgb = NAVY if alert_type == "NOTE" else RGBColor(180, 50, 50)
    
    body = p.add_run(text)
    body.font.italic = True
    body.font.color.rgb = CHARCOAL
