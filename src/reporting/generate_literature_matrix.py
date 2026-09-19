import pandas as pd
import os
import shutil
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

def generate_literature_matrix():
    # Base directory
    base_dir = r"c:\Users\ARYAN - AYUSH\OneDrive\Desktop\an intership\pharmaceutical_research_portfolio"
    out_dir = os.path.join(base_dir, "week1_literature_review")
    ref_dir = os.path.join(out_dir, "references")
    master_ref = os.path.join(base_dir, "references", "master_references.bib")
    
    os.makedirs(out_dir, exist_ok=True)
    os.makedirs(ref_dir, exist_ok=True)
    
    # Data
    columns = [
        "ID", "Title", "Authors", "Year", "Journal", "DOI", "PMID", 
        "Research_Objective", "Study_Design", "Population_Sample", 
        "Methodology", "Major_Findings", "Limitations", "Relevance", 
        "Research_Gap_Identified"
    ]
    
    data = [
        ["1", "Adverse events in different administration routes of semaglutide: FAERS study", "Author", "2024", "Frontiers in Pharmacology", "10.3389/fphar.2024.1414268", "38887555", "Evaluate ADRs of semaglutide by route", "Retrospective pharmacovigilance", "FAERS database", "Disproportionality analysis", "Different ADR profiles for oral vs subcutaneous routes", "Underreporting bias", "High", "Specific ADR differences by route"],
        ["2", "Comparative analysis semaglutide FAERS + social media", "Author", "2024", "Frontiers in Pharmacology", "10.3389/fphar.2024.1471615", "39502525", "Compare FAERS and social media ADRs", "Mixed-methods pharmacovigilance", "FAERS and social media users", "Data mining and disproportionality", "Social media captures ADRs not in FAERS", "Selection bias", "High", "Value of social media in pharmacovigilance"],
        ["3", "Real-world disproportionality analysis semaglutide", "Author", "2024", "J Diabetes Investig", "10.1111/jdi.14229", "38943656", "Analyze real-world semaglutide ADRs", "Retrospective pharmacovigilance", "FAERS database", "Disproportionality analysis", "Multiple disproportionate signals across organ systems", "Causality not established", "High", "Comprehensive overview of real-world ADRs"],
        ["4", "Semaglutide NAION FAERS", "Author", "2025", "Obesity Res Clin Pract", "10.1016/j.orcp.2025.01.011", "39922760", "Investigate NAION risk", "Retrospective pharmacovigilance", "FAERS database", "Disproportionality analysis", "Significant NAION signal", "Confounding by indication", "High", "Rare optic nerve ADRs"],
        ["5", "Mortality and SAEs with GLP-1 RAs in FAERS", "Author", "2024", "Cureus", "10.7759/cureus.65989", "39221363", "Evaluate serious ADRs and mortality", "Retrospective pharmacovigilance", "FAERS database", "Disproportionality analysis", "Patterns of serious adverse events", "No incidence rate calculation", "High", "Severe outcomes with GLP-1 RAs"],
        ["6", "Depression/suicide signals for weight loss meds in FAERS", "Author", "2025", "J Affect Disord", "10.1016/j.jad.2025.119670", "40523410", "Assess psychiatric safety", "Retrospective pharmacovigilance", "FAERS database", "Disproportionality analysis", "Depression and suicidality signals identified", "Weber effect", "High", "Psychiatric ADRs for weight loss"],
        ["7", "READUS-PV reporting checklist", "Author", "2024", "Drug Safety", "10.1007/s40264-024-01421-9", "38713347", "Develop reporting guidelines for PV", "Methodological guideline", "PV literature", "Consensus methodology", "14-item reporting checklist for disproportionality analyses", "Requires broad adoption", "Medium", "Standardization of PV reporting"],
        ["8", "STEP 1 Trial - Wilding et al.", "Wilding et al.", "2021", "NEJM", "10.1056/NEJMoa2032183", "33567185", "Evaluate semaglutide for obesity", "RCT Phase 3", "Adults with overweight/obesity", "Double-blind, placebo-controlled", "14.9% weight reduction, GI adverse events", "Strict inclusion criteria", "High", "Base efficacy/safety data"],
        ["9", "WHO data semaglutide suicidality", "Author", "2024", "JAMA Network Open", "10.1001/jamanetworkopen.2024.23385", "39163046", "Analyze global suicidality reports", "Retrospective pharmacovigilance (WHO)", "VigiBase", "Disproportionality analysis", "Suicidal ideation signal with antidepressant co-prescription", "Database overlap", "High", "Global perspective on psychiatric ADRs"],
        ["10", "GLP-1 RAs suicidality FAERS", "Author", "2024", "Expert Opin Drug Saf", "10.1080/14740338.2023.2295397", "38087976", "Assess causality of suicidal signals", "Retrospective pharmacovigilance", "FAERS database", "Disproportionality analysis + Bradford Hill", "Disproportionate reporting but no causal link (Bradford Hill)", "Retrospective nature", "High", "Causality assessment of psych ADRs"],
        ["11", "Semaglutide suicidal ideation real-world cohort", "Author", "2024", "Nature Medicine", "10.1038/s41591-023-02672-2", "38182782", "Compare suicidal ideation incidence", "Retrospective cohort", "EHR data", "Propensity score matching", "LOWER risk (contradicts FAERS)", "Residual confounding", "High", "Real-world incidence vs spontaneous reporting"],
        ["12", "GI safety GLP-1 RAs FAERS", "Author", "2024", "Diagnostics", "", "39767190", "Evaluate gastrointestinal safety", "Retrospective pharmacovigilance", "FAERS database", "Disproportionality analysis", "Strong GI ADR association", "Known ADR bias", "High", "Detailed GI profile"],
        ["13", "Anti-obesity meds digestive ADEs FAERS", "Author", "2024", "BMC Pharmacol Toxicol", "", "39267168", "Analyze digestive ADEs for anti-obesity drugs", "Retrospective pharmacovigilance", "FAERS database", "Disproportionality analysis", "Pancreatitis signals", "Multiple comparisons", "High", "Comparative digestive safety"],
        ["14", "GLP-1 RA metabolic/nutritional ADEs", "Author", "2024", "Front Pharmacol", "", "39040467", "Assess metabolic and nutritional ADEs", "Retrospective pharmacovigilance", "FAERS database", "Disproportionality analysis", "Dehydration and hypoglycemia profiles", "Missing clinical context", "High", "Metabolic consequences of treatment"],
        ["15", "Ophthalmic ADRs GLP-1 RAs FAERS", "Author", "2025", "Endocrine", "", "39578328", "Investigate ophthalmic safety", "Retrospective pharmacovigilance", "FAERS database", "Disproportionality analysis", "Retinopathy and visual impairment signals", "Diabetic retinopathy confounding", "High", "Ophthalmic ADR spectrum"],
        ["16", "Alopecia semaglutide tirzepatide FAERS", "Author", "2024", "JEADV", "", "38925559", "Evaluate hair loss risk", "Retrospective pharmacovigilance", "FAERS database", "Disproportionality analysis", "Significant alopecia signal", "Cosmetic concern underreporting", "Medium", "Dermatological ADRs"],
        ["17", "Alopecia GLP-1 RAs scoping review", "Author", "2025", "JAAD Case Reports", "", "40787040", "Review clinical evidence for alopecia", "Scoping review", "Published case reports", "Literature review", "Clinical cases corroborate FAERS alopecia signal", "Publication bias", "Medium", "Clinical validation of FAERS signal"]
    ]
    
    df = pd.DataFrame(data, columns=columns)
    
    # Save CSV
    csv_path = os.path.join(out_dir, "literature_matrix.csv")
    df.to_csv(csv_path, index=False)
    
    # Save Excel
    xlsx_path = os.path.join(out_dir, "literature_matrix.xlsx")
    writer = pd.ExcelWriter(xlsx_path, engine='openpyxl')
    df.to_excel(writer, index=False, sheet_name='Literature Matrix')
    
    workbook = writer.book
    worksheet = writer.sheets['Literature Matrix']
    
    # Formatting
    header_font = Font(bold=True)
    wrap_alignment = Alignment(wrap_text=True, vertical='top')
    
    for cell in worksheet["1:1"]:
        cell.font = header_font
    
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap_alignment
            
    # Column widths
    for i, col in enumerate(columns):
        col_letter = get_column_letter(i+1)
        if col in ["Title", "Research_Objective", "Major_Findings", "Limitations", "Research_Gap_Identified"]:
            worksheet.column_dimensions[col_letter].width = 40
        else:
            worksheet.column_dimensions[col_letter].width = 20
            
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.freeze_panes = "A2"
    
    writer.close()
    
    # References
    df[['ID', 'PMID', 'DOI', 'Title', 'Authors', 'Year', 'Journal']].to_csv(os.path.join(ref_dir, "references.csv"), index=False)
    
    if os.path.exists(master_ref):
        shutil.copy(master_ref, os.path.join(ref_dir, "references.bib"))
    else:
        with open(os.path.join(ref_dir, "references.bib"), 'w') as f:
            f.write("%% Master references not found, creating empty bib\n")
            
    print(f"Generated files in {out_dir}")

if __name__ == "__main__":
    generate_literature_matrix()
